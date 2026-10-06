from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "ppt"
WIDTH = 13.333
HEIGHT = 7.5
INK = RGBColor(28, 43, 52)
MUTED = RGBColor(81, 101, 108)
PAPER = RGBColor(247, 249, 246)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(16, 113, 108)
CORAL = RGBColor(224, 111, 82)
PALE_TEAL = RGBColor(226, 241, 237)
PALE_CORAL = RGBColor(250, 235, 228)


# Each lesson is a group of precise curriculum topics. Every topic has a
# short explanation and an example suitable for a live classroom walkthrough.
MODULES = [
    {
        "id": "01", "title": "Programming Fundamentals & Web Development Basics",
        "duration": "5 hours", "outcome": "Explain a web request and express a small real-world rule as testable logic.",
        "thread": "A student registers for a campus event; an organizer tracks a small event budget.",
        "lessons": [
            ("Software development in practice", [
                ("Software development", "A repeatable process for understanding a need, building a solution, checking it, and improving it.", "For a festival signup, first ask who registers and what confirmation they need."),
                ("SDLC", "The software development life cycle commonly moves through requirements, design, implementation, testing, release, and maintenance.", "After releasing the event page, fix a date bug and verify the deployed page: maintenance is still development."),
                ("Types of applications", "Web, mobile, desktop, and embedded applications differ by where they run and how users access them.", "A browser-based event portal is a web app; a campus kiosk may be a desktop or embedded app."),
                ("Team roles", "Product, design, frontend, backend, database, QA, and operations roles contribute different expertise; small teams may share roles.", "A QA partner checks that a full event cannot accept another registration."),
                ("Full-stack career path", "A full-stack developer can work across UI, server, and data boundaries while still seeking review in unfamiliar areas.", "Trace one registration through React, Express, and MongoDB before choosing which layer to improve."),
            ]),
            ("How the web communicates", [
                ("Internet fundamentals", "The internet is a network of connected devices; the web is one service that uses it to exchange linked resources.", "A browser uses a network connection to request an event page from a host."),
                ("Client-server architecture", "A client requests a resource or action; a server receives it, applies rules, and returns a response.", "The browser asks the event server whether seats remain; the server answers with availability."),
                ("Frontend, backend, database", "The frontend presents interaction, backend enforces application rules, and database persists records.", "A form collects an event choice; an API checks capacity; a database stores the registration."),
                ("HTTP request-response", "An HTTP request carries a method, URL, headers, and sometimes a body; the response carries a status, headers, and data.", "`POST /api/registrations` with an event ID can return `201 Created` and a confirmation record."),
            ]),
            ("Values and expressions", [
                ("Variables", "A variable names a value so a program can use or update it; choose names that describe meaning.", "`const remainingSeats = 12` is clearer than `const x = 12`."),
                ("Data types", "Types describe the kind of value, such as number, string, boolean, array, or object; operations depend on type.", "`" + "\"Robotics Meetup\"" + "` is text, `12` is a number, and `isOpen` can be boolean."),
                ("Operators", "Arithmetic, comparison, and logical operators combine or compare values to produce a result.", "`remainingSeats > 0 && isPublished` checks two conditions before offering registration."),
            ]),
            ("Control flow and reusable logic", [
                ("Conditionals", "An `if/else` selects a path based on a true/false condition.", "If seats are greater than zero, accept the request; otherwise return “Event full”."),
                ("Loops", "A loop repeats work for each item or until a condition changes; a clear stopping condition prevents infinite work.", "Loop over three expense amounts and add each to a running total."),
                ("Functions", "A function packages a named operation with inputs and an optional returned result.", "`calculateTotal(prices)` returns the sum so multiple reports can reuse the rule."),
                ("Arrays and objects", "An array stores an ordered list; an object groups named properties that describe one entity.", "`[{ title: 'Talk', seats: 20 }]` is an array of event objects."),
            ]),
            ("Problem solving before code", [
                ("Algorithm design", "An algorithm is a finite, ordered set of steps that solves a defined problem.", "Validate an event ID, look up the event, check capacity, save registration, then confirm."),
                ("Flowcharts", "A flowchart uses standard shapes and arrows to show steps and decision branches.", "A diamond asks `seats > 0?`; Yes leads to save, No leads to a full-event message."),
                ("Pseudocode", "Pseudocode describes logic in readable, language-neutral steps before syntax is chosen.", "`IF seats > 0 THEN confirm ELSE report full` communicates the rule without JavaScript details."),
                ("Break down real problems", "Translate a broad request into users, inputs, rules, outputs, and edge cases.", "For expense tracking, define whether refunds count before writing a savings calculator."),
            ]),
        ],
        "practice": "Complete the paired request-map and expense-summary lab. Test available seats, a full event, and a negative/invalid capacity before accepting the logic.",
    },
    {
        "id": "02", "title": "HTML, CSS & JavaScript Fundamentals",
        "duration": "10 hours", "outcome": "Build an accessible, responsive event page with a form that gives useful feedback.",
        "thread": "Present one fictional campus workshop clearly on a phone and desktop.",
        "lessons": [
            ("HTML structure and meaning", [
                ("HTML5 and semantic HTML", "HTML describes document structure; semantic elements communicate the role of content to browsers and assistive technology.", "Use `<main>`, `<article>`, `<h1>`, and `<footer>` for the event page instead of generic containers everywhere."),
                ("Forms and input controls", "A form groups user input; each control needs an appropriate type, name, and associated label.", "Use `<label for='email'>Email</label>` with `<input id='email' type='email'>`."),
                ("Tables", "Tables represent tabular data with row/column relationships, not general page layout.", "A schedule with session, time, and room is tabular; use `<th>` for column headings."),
                ("Multimedia", "Audio and video need controls and accessible alternatives; images need meaningful alternative text when informative.", "A festival photo gets a concise alt description; a decorative flourish gets empty alt text."),
                ("Accessibility basics", "Accessible pages support keyboard use, meaningful labels, readable contrast, and understandable feedback.", "Tab to Register and confirm a visible focus indicator appears."),
            ]),
            ("CSS layout and responsive design", [
                ("Selectors", "CSS selectors target elements by type, class, or relationship; specific, reusable classes are easier to maintain.", "`.event-card` styles every event card without repeating inline styles."),
                ("Box model", "An element's rendered size includes content, padding, border, and margin; `box-sizing` affects the calculation.", "A 300px-wide card with padding can overflow unless sizing accounts for padding."),
                ("Flexbox", "Flexbox arranges items along one primary axis and distributes available space.", "Place event date and Register button in one aligned card row."),
                ("Grid layout", "CSS Grid defines rows and columns for two-dimensional layouts.", "Display event cards in a three-column desktop grid that becomes one column on mobile."),
                ("Responsive design and media queries", "Responsive design adapts layout to available space; media queries apply styles when conditions match.", "At a narrow viewport, change the card grid from three columns to one."),
                ("Animations", "CSS animation can communicate state or draw attention, but should be restrained and respect reduced-motion preferences.", "A short confirmation fade-in is optional; the success message must remain understandable without motion."),
            ]),
            ("JavaScript fundamentals", [
                ("Variables and data types", "Use `const` for bindings that are not reassigned and `let` when reassignment is needed; values retain their types.", "`const eventTitle = 'Design Jam'` names text displayed in the card."),
                ("Functions", "Functions isolate behavior so it can be called, tested, and reused.", "`showMessage(text)` updates the same status area after submit or validation."),
                ("Arrays and objects", "Arrays hold lists and objects represent named fields; these structures model repeated page data.", "Map an array of `{ id, title, venue }` objects into event cards."),
                ("DOM manipulation", "The Document Object Model exposes a page as nodes that JavaScript can read and update.", "Set a status element's `textContent` after a successful form submission."),
                ("Event handling", "An event listener runs a function in response to user or browser activity.", "Listen for the form's `submit` event rather than relying on a button click alone."),
                ("Form validation", "Validation checks user input before processing; browser validation improves UX, while the server must enforce important rules too.", "`required` and `type='email'` catch obvious mistakes before sending a registration."),
            ]),
        ],
        "practice": "Build the accessible event card and registration form. Check valid/invalid submission, keyboard navigation, and two viewport widths.",
    },
    {
        "id": "03", "title": "Advanced JavaScript (ES6+) & Git/GitHub",
        "duration": "15 hours", "outcome": "Load remote data safely and collaborate through small, reviewable Git changes.",
        "thread": "Fetch event information, show network states, and submit a reviewed feature branch.",
        "lessons": [
            ("Modern JavaScript syntax", [
                ("`let` and `const`", "`const` prevents rebinding; `let` allows rebinding. Both are block-scoped and usually preferred over `var`.", "Keep `const events = []`; use `let page = 1` only if page changes."),
                ("Arrow functions", "Arrow functions provide concise function syntax and inherit `this` from their surrounding scope.", "`events.filter((event) => event.isOpen)` keeps the filtering rule readable."),
                ("Template literals", "Backticks allow interpolation and multi-line strings without manual concatenation.", "`` `Hello, ${studentName}` `` creates a personalized local greeting."),
                ("Destructuring", "Destructuring extracts named or indexed values from objects and arrays.", "`const { title, venue } = event` names the two fields a card needs."),
                ("Spread and rest", "Spread expands iterable/object values; rest gathers remaining values into one array or object.", "`const next = [...events, newEvent]` creates a new array without mutating the old one."),
                ("Modules", "Modules split code into files with explicit exports and imports, making dependencies easier to see.", "Export `fetchEvents` from `eventApi.js` and import it in the page component."),
            ]),
            ("Asynchronous code and HTTP", [
                ("Callbacks", "A callback is a function passed to another operation to run later; deeply nested callbacks can obscure control flow.", "A timer callback runs after a delay; network APIs historically used callbacks for completion."),
                ("Promises", "A Promise represents a future result and can be pending, fulfilled, or rejected.", "`fetch(url).then(readResponse).catch(showError)` chains response handling."),
                ("Async/await", "`async` functions return Promises; `await` makes Promise-based code read sequentially while remaining asynchronous.", "`const response = await fetch(url)` waits for the response before parsing JSON."),
                ("Fetch API", "`fetch` sends HTTP requests; callers must check `response.ok` because HTTP error statuses do not automatically reject the Promise.", "A 404 can still resolve, so check status before rendering event data."),
            ]),
            ("Debugging and reliable failures", [
                ("Error handling", "Handle expected failures at the boundary and show useful user feedback without exposing internal details.", "A failed event fetch displays Retry instead of an empty list that looks successful."),
                ("Debugging tools", "Console, debugger, and network panels help inspect values, call stacks, requests, and response status.", "Inspect a failed request's URL and status before changing the rendering code."),
                ("Try/catch", "`try/catch` handles exceptions from awaited operations and keeps failure behavior explicit.", "Catch JSON/network errors and set an error state for the event list."),
                ("Common JavaScript errors", "Syntax, reference, type, and async errors have different causes; use the first relevant stack frame and reproduce reliably.", "A misspelled `event.titel` yields undefined; inspect the object shape and correct the field name."),
            ]),
            ("Git and GitHub collaboration", [
                ("Version control", "Version control records project changes over time so developers can compare, restore, and collaborate.", "A commit shows exactly when the event filter was added."),
                ("Repository management", "A repository holds tracked files and history; `.gitignore` excludes generated files and local secrets.", "Commit source, not `node_modules` or `.env`."),
                ("Branching", "A branch isolates a line of work so a feature can be developed without destabilizing the shared base.", "Create `feature/event-search` from the current main branch."),
                ("Pull requests", "A pull request proposes a branch for review, discussion, checks, and eventual integration.", "Describe the search behavior, tests, and known limitation in the PR."),
                ("Merge conflicts", "A merge conflict occurs when Git cannot reconcile overlapping edits; a person chooses the intended combined result.", "Keep both a teammate's venue field and your search logic, then run the page tests."),
            ]),
        ],
        "practice": "Implement an event feed with loading, error, and empty states on a feature branch; request peer review and verify the diff.",
    },
    {
        "id": "04", "title": "React.js Frontend Development",
        "duration": "25 hours", "outcome": "Build a reusable React event browser with controlled state, navigation, and API states.",
        "thread": "Compose a student-facing event browser from reusable UI and data-flow boundaries.",
        "lessons": [
            ("React mental model", [
                ("React introduction", "React builds user interfaces from components that render from data and update when state changes.", "Changing a search query rerenders the visible event results."),
                ("Virtual DOM", "React describes desired UI output and reconciles changes to update the browser DOM; it is a programming model, not a guarantee every app is faster.", "Changing one registration status updates the relevant rendered UI from new state."),
                ("JSX", "JSX is syntax for describing UI in JavaScript; expressions use braces and components use capitalized names.", "`<EventCard title={event.title} />` passes a value into a component."),
                ("Project structure and app creation", "A React project separates entry point, components, pages, assets, and services; scaffolding provides build and dev tooling.", "Place shared event API calls in a service instead of embedding URLs in every card."),
            ]),
            ("Components, props and composition", [
                ("Functional components", "A component is a function that returns UI for its inputs and state.", "`function EventCard({ event }) { return <article>{event.title}</article> }`."),
                ("Reusable components", "Reuse a component when the same presentation/behavior appears with different data; avoid abstractions that hide simple cases.", "Render 20 events with one `EventCard` implementation."),
                ("Props and data flow", "Props are read-only inputs passed from a parent to a child; data generally flows down.", "The list passes each event and an `onRegister` handler to its card."),
                ("Component composition", "Composition combines small components through nesting and children rather than building one oversized component.", "An `EventPage` composes `SearchBar`, `EventList`, and `RegistrationDialog`."),
            ]),
            ("State and hooks", [
                ("State versus props", "Props come from a parent; state is data a component owns and updates over time.", "An event title is a prop; a local open/closed dialog flag is state."),
                ("`useState`", "`useState` stores local component state and returns its current value and setter.", "Store the current search string with `const [query, setQuery] = useState('')`."),
                ("State updates", "State setters schedule a render; when next state depends on previous state, use the functional updater form.", "`setCount((count) => count + 1)` avoids relying on a stale count."),
                ("State lifting", "Lift shared state to the nearest common ancestor when multiple components need to read or update it.", "Keep search text in the page if both toolbar and event list depend on it."),
                ("`useEffect`", "Effects synchronize with external systems after render; dependencies describe which reactive values the effect uses.", "Fetch events when the page mounts or the selected category changes."),
                ("`useRef`", "A ref stores a mutable value that does not trigger rendering; it can also reference a DOM element.", "Focus the email input after opening the registration dialog."),
                ("Custom hooks", "A custom hook extracts reusable stateful logic while following the Rules of Hooks.", "`useEvents(category)` can own loading, data, and error state for event fetching."),
            ]),
            ("Interactions, forms and navigation", [
                ("Event handling", "Pass a function to an event prop; do not call it during render unless the call is intentionally producing a value.", "`onClick={() => onRegister(event.id)}` registers only after the click."),
                ("Controlled forms", "A controlled input takes its displayed value from React state and updates it through an event handler.", "`value={email}` and `onChange` keep the current form value explicit."),
                ("Form validation and dynamic forms", "Validate required fields and relationships; derive repeated controls from data while preserving stable keys.", "Show one attendee field per selected ticket and reject an empty email."),
                ("React Router and navigation", "Client-side routing maps URLs to views without a full document reload; links preserve browser navigation behavior.", "`/events/42` renders details for event 42."),
                ("Route parameters", "Route parameters are URL segments that identify a resource and should be validated before use.", "Read `eventId` from `/events/:eventId`, then request the matching record."),
                ("Protected routes basics", "A protected UI route improves navigation but server authorization must still protect private data/actions.", "Hide organizer tools for students, and separately enforce organizer role in the API."),
            ]),
            ("APIs, shared state and quality", [
                ("Fetch and Axios", "Fetch is built into browsers; Axios is a library with additional request/response conveniences. Both need clear error handling.", "A service function requests `/api/events` and returns parsed event data."),
                ("GET, POST, PUT, DELETE", "HTTP methods express resource intent: read, create, replace/update, and remove; APIs should define consistent contracts.", "Use GET for events and POST for a new registration."),
                ("API errors", "Represent loading, success, empty, and error as distinct states so a failure is not mistaken for no data.", "Show “Could not load events” on network failure and preserve a retry action."),
                ("Context API", "Context shares values across a subtree without passing props through every intermediate component; overuse hides dependencies.", "Share current user identity with navigation and profile views."),
                ("Lazy loading and code splitting", "Lazy loading defers loading a component until needed; code splitting divides bundles into smaller chunks.", "Load a rarely used organizer report when its route is visited."),
                ("Memoization concepts", "Memoization can skip repeated work when inputs are unchanged; measure first because it adds complexity.", "Profile a large event filter before memoizing its expensive calculation."),
                ("Responsive components and UI patterns", "Components should work at available widths and expose consistent states, semantics, and interaction patterns.", "An event card stacks date above title on a narrow screen and keeps actions keyboard-accessible."),
            ]),
        ],
        "practice": "Complete the searchable event browser with components, form feedback, and optional event routes/API integration; explain one state ownership decision.",
    },
    {
        "id": "05", "title": "Backend Development with Node.js & Express.js",
        "duration": "20 hours", "outcome": "Design, implement, and test a small REST API with validation and consistent errors.",
        "thread": "Build an event API that a frontend can call and an API client can verify.",
        "lessons": [
            ("Backend and Node.js", [
                ("Backend architecture", "A backend receives requests, applies business rules, coordinates persistence, and returns a contract to clients.", "Check event capacity on the server before writing a registration."),
                ("Client-server communication", "Clients and servers communicate through explicit request/response contracts over a network.", "A frontend sends JSON to `/api/events`; Express returns a status and JSON body."),
                ("REST principles and APIs", "REST-style APIs expose resources through URLs and standard HTTP semantics; an API is a contract between systems.", "`GET /api/events/42` reads event 42 rather than using an action name in the URL."),
                ("Node.js", "Node.js runs JavaScript outside the browser and provides APIs for servers, files, and networking.", "Run `node server.js` to start an Express API process."),
                ("Event-driven architecture", "Node handles many I/O operations through an event loop rather than blocking one thread for each request.", "A server can wait for a database response while handling other incoming work."),
                ("Modules, packages and npm", "Modules split source; packages provide reusable code; npm manages dependencies and project scripts.", "Install Express as a dependency and import it from the server entry point."),
                ("File system operations", "Node's file-system APIs read and write files; asynchronous operations avoid blocking the event loop for routine I/O.", "Read a local seed JSON file for a classroom demo, not as a substitute for a production database."),
            ]),
            ("Express request pipeline", [
                ("Express application", "An Express app configures middleware and routes, then listens for HTTP requests on a port.", "`app.listen(3000)` starts the local event API."),
                ("Routing", "Routes match an HTTP method and path to the code responsible for that resource/action.", "`app.get('/api/events', handler)` handles event-list requests."),
                ("Middleware", "Middleware runs in order, can inspect/change request or response, and must continue, respond, or pass an error onward.", "`express.json()` parses a JSON request body before the route reads it."),
                ("Request and response handling", "Handlers read validated route/query/body data and produce a status plus a response body.", "Return `201` and the created event after valid input is stored."),
            ]),
            ("REST CRUD and API contracts", [
                ("GET, POST, PUT, DELETE", "CRUD APIs use methods consistently: read, create, replace/update, and delete a resource.", "GET the list, POST a new event, PATCH its venue, DELETE it by ID."),
                ("API design", "Good API contracts use predictable resource paths, status codes, field names, and error shapes.", "Return `{ error: 'title is required' }` with 400 for a missing title."),
                ("Request validation", "Validate untrusted input at the server boundary for shape, type, allowed values, and business constraints.", "Reject a start date in the past or a negative seat count before database writes."),
                ("Error handling", "Centralized error handling avoids duplicated response logic and prevents internal details leaking to clients.", "Log diagnostic context privately; return a generic 500 message to the client."),
            ]),
            ("Middleware, testing and security", [
                ("Custom and logging middleware", "Custom middleware implements cross-cutting request behavior such as timing or request IDs.", "Log method, path, and duration without logging passwords or tokens."),
                ("Authentication middleware concepts", "Authentication middleware verifies a credential and attaches trusted identity; authorization is a separate permission decision.", "Verify a token before an organizer-only event update route."),
                ("Postman and API testing", "An API client sends controlled requests and lets teams inspect status, headers, body, and repeatable collections.", "Test create with valid JSON and then repeat with the required title omitted."),
                ("API debugging", "Reproduce a failure, inspect request/response and server logs, then isolate the smallest failing boundary.", "A 404 may be a route mismatch; inspect URL and method before changing database code."),
                ("Environment variables", "Environment variables configure per-environment values without hard-coding secrets into source.", "Read a database URL from `process.env` and keep the local `.env` ignored."),
                ("Input sanitization and security", "Validation constrains acceptable input; escaping/parameterization prevents interpretation as code; least privilege limits impact.", "Use a database query API rather than concatenating user input into a query string."),
            ]),
        ],
        "practice": "Implement event CRUD routes and validation. Test successful and invalid requests, unknown IDs, status codes, and error messages in Postman or curl.",
    },
    {
        "id": "06", "title": "MongoDB Database Development",
        "duration": "10 hours", "outcome": "Model event data, perform CRUD/query operations, and explain a useful index or aggregation.",
        "thread": "Persist events and answer organizer questions about categories and remaining stock/seats.",
        "lessons": [
            ("Database foundations and MongoDB", [
                ("Data storage concepts", "A database persists structured information so it remains available beyond one running process.", "Event records remain after the API server restarts."),
                ("SQL vs NoSQL", "Relational databases organize related tables with schemas and joins; document databases store flexible documents and related collections.", "An event document can keep venue details together; analytics may still use a relational model well."),
                ("Database design principles", "Model around data ownership, consistency rules, expected reads/writes, and data lifecycle—not just fields on a screen.", "Keep registrations separate if they grow independently from an event's details."),
                ("MongoDB architecture", "MongoDB stores BSON documents in collections inside databases and supports queries and indexes over fields.", "The `campus_events` database contains an `events` collection."),
                ("Collections and documents", "A collection groups documents; each document stores named fields and has a unique `_id`.", "`{ _id, title, venue, seats }` is one event document."),
                ("BSON", "BSON is MongoDB's binary-encoded document format with types beyond plain JSON, including ObjectId and dates.", "Store `startsAt` as a date value rather than a formatted display string."),
            ]),
            ("CRUD and query patterns", [
                ("Insert and read", "Insert creates documents; read queries select documents matching a filter.", "Insert a workshop, then find all events whose status is `open`."),
                ("Update and delete", "Update modifies selected fields/documents; delete removes selected documents, so filters should be deliberate.", "Update seats by event `_id`; verify the changed document with a follow-up read."),
                ("Filters and query operators", "Filters match field values; operators express comparisons and logical conditions.", "Find events with `seats: { $gt: 0 }` to list those with capacity."),
                ("Sorting and pagination", "Sort produces predictable order; pagination limits result size and uses a stable order to avoid duplicates/skips.", "Sort by start date then `_id`, and request 10 events after a cursor."),
            ]),
            ("Aggregation, Mongoose and performance", [
                ("Aggregation framework", "An aggregation pipeline transforms documents through ordered stages such as match, group, project, and sort.", "Match published events, group by category, and count each group."),
                ("Grouping and calculations", "Group stage collects documents by a key and computes accumulators such as sum, average, or count.", "Group registrations by event and sum the number of seats reserved."),
                ("Mongoose ODM", "Mongoose maps application objects to MongoDB documents and provides schemas, models, validation, and middleware.", "Use `Event.find({ status: 'open' })` through the Event model."),
                ("Schemas and models", "A schema describes expected fields and validation; a model provides operations for that document type.", "Require a title and restrict seats to a nonnegative number in the Event schema."),
                ("Relationships", "References connect independently managed documents; embedding can be simpler when related data is small and read together.", "Reference the event ID from each registration rather than copying full event details."),
                ("Indexing and query optimization", "An index can speed matching/sorting at the cost of storage and slower writes; create it for measured query patterns.", "Index `status` and `startsAt` if the common open-events query filters and sorts on them."),
            ]),
        ],
        "practice": "Create the event schema and query set, including sorted open events and a grouped category report. Explain why each field or index exists.",
    },
    {
        "id": "07", "title": "Full Stack Integration, APIs & Authentication",
        "duration": "15 hours", "outcome": "Trace a complete client-to-database workflow and protect it with authentication and authorization.",
        "thread": "Register an attendee for an event and restrict organizer actions to permitted users.",
        "lessons": [
            ("MERN architecture and integration", [
                ("MERN stack architecture", "React renders the client, Express/Node handles API rules, and MongoDB persists data; each boundary has a contract.", "React calls an Express route that reads/writes an Event model in MongoDB."),
                ("Frontend-backend communication", "A client service sends HTTP requests and translates response states into user-visible behavior.", "A registration button sends a POST and displays confirmation only after success."),
                ("Data flow across layers", "Data is validated at trust boundaries, transformed as needed, persisted, and returned with a defined shape.", "UI form values become a validated API body, a database record, then a confirmation response."),
                ("Connect React, Express and database", "Configure the client API base URL, server routes, and database connection per environment.", "Local React at one port calls the API at another, which uses a local or sandbox MongoDB."),
                ("End-to-end CRUD", "End-to-end operations prove that UI, API, and persistence work together for a user task.", "Create an event, reload the browser, and confirm it still appears from the database."),
                ("Reusable services and folder structure", "Separating UI, API, models, middleware, and configuration makes responsibilities easier to locate and test.", "Keep `eventApi` calls outside the visual `EventCard` component."),
            ]),
            ("Authentication and authorization", [
                ("Registration and login", "Registration creates an identity; login verifies credentials and establishes an authenticated session/token.", "A test student account logs in before viewing their own registrations."),
                ("Password hashing", "Store a salted password hash using a maintained password-hashing library; never store or log plaintext passwords.", "On login, compare the submitted password with the stored hash using the library."),
                ("JWT authentication", "A signed JWT conveys claims that a server verifies; signing does not encrypt its payload, and expiry/secret handling matter.", "A token can identify a user ID and expiry, but should not contain a password or sensitive profile."),
                ("Role-based access control", "RBAC grants actions based on roles and resource rules; the server must enforce permission on each protected action.", "An organizer may edit event details while a student may only register."),
                ("Protected routes and permissions", "A protected UI route guides navigation; backend authorization is the security boundary for data/actions.", "Directly call the organizer endpoint as a student and verify a 403 response."),
            ]),
            ("Files, configuration and security", [
                ("File uploads and image storage", "Uploads require limits, type validation, safe names, and a storage policy; avoid trusting the client filename or MIME claim alone.", "Allow only bounded image uploads for an event poster and store outside the application source tree."),
                ("Profile management", "Profile APIs should expose only necessary fields and enforce ownership or role checks.", "A student can update their own display name but cannot change their account role."),
                ("Security best practices", "Use least privilege, validate inputs, protect secrets, limit sensitive data, and keep dependencies maintained.", "Do not return password hashes in a registration response."),
                ("Environment configuration", "Separate development/test/production config and inject secrets at runtime through the deployment environment.", "Read the token signing secret from configuration; never commit its value."),
            ]),
        ],
        "practice": "Connect the event UI, API, and database; test successful registration and an organizer-only action as both permitted and denied users.",
    },
    {
        "id": "08", "title": "Deployment, Testing & DevOps Fundamentals",
        "duration": "10 hours", "outcome": "Test key behaviors, configure environments, and verify a release from build to live smoke test.",
        "thread": "Release the event application safely and diagnose a failed registration after deployment.",
        "lessons": [
            ("Testing fundamentals", [
                ("Importance of testing", "Tests provide repeatable evidence that important behavior still works; they reduce risk but cannot prove absence of all defects.", "A regression test prevents a full event from accepting another attendee."),
                ("Unit, integration and functional tests", "Unit tests check a small unit; integration tests check collaborating parts; functional tests validate user-visible behavior.", "Unit-test a capacity rule, integration-test its API/database path, then test registration in the browser."),
                ("API testing", "API tests assert request/response contracts including status, body, and important failure cases.", "POST valid registration expects 201; POST without an event ID expects 400."),
                ("Component and UI tests", "Component tests verify rendered behavior and interaction; UI validation checks that users can perceive and operate the interface.", "Submitting an invalid form shows an associated error message."),
                ("Request validation and error tests", "Failure-path tests ensure invalid input and unexpected conditions produce safe, useful outcomes.", "An unknown event ID returns 404 and does not create a registration."),
            ]),
            ("Build, configuration and hosting", [
                ("Build process", "A build transforms source into artifacts suited to production; successful compilation alone does not verify user workflows.", "Build the React bundle, then separately smoke-test event listing."),
                ("Environment variables", "Environment-specific values are injected at runtime and should not be committed or printed in logs.", "Configure API URL and database connection in the host's secret/config settings."),
                ("Production configuration", "Production settings enable safe error handling, correct origins, secure cookies/HTTPS, and appropriate resource limits.", "Do not use a development wildcard CORS policy for a public production API."),
                ("Frontend, backend and database hosting", "A full-stack release may involve separately hosted client, API, and database services with network permissions between them.", "Confirm the hosted frontend can reach the API and the API can reach the database."),
                ("Domain and SSL", "DNS maps a domain to a service; TLS certificates encrypt and authenticate HTTPS connections.", "Check the deployed event URL loads over HTTPS without certificate warnings."),
            ]),
            ("DevOps and operations", [
                ("DevOps", "DevOps combines development and operations practices to deliver changes reliably through automation, feedback, and shared ownership.", "Developers own a tested release and inspect production logs after deployment."),
                ("CI/CD", "Continuous integration runs automated checks on changes; continuous delivery/deployment prepares or releases validated changes.", "A pull request triggers install, tests, and build before merge."),
                ("GitHub Actions pipeline", "A workflow file defines automated jobs triggered by repository events; credentials should be scoped and protected.", "Run `npm test` and `npm run build` on each pull request without deploying secrets to untrusted forks."),
                ("Logging and monitoring", "Logs record events for diagnosis; monitoring tracks service health and performance over time.", "A request ID links a sanitized API error log to a failing registration report."),
                ("Error tracking and performance", "Error tracking groups failures; performance measures help identify slow routes or resources worth improving.", "Measure event-list response time before deciding that an index is required."),
            ]),
        ],
        "practice": "Add tests for one success and one failure, build for production, then complete the deployment smoke-test and recovery checklist.",
    },
    {
        "id": "09", "title": "Capstone Project Studio",
        "duration": "24 hours", "outcome": "Plan, build, test, deploy, and present a complete MVP that solves a defined user problem.",
        "thread": "Teams choose a README project idea and deliver a demonstrable end-to-end slice.",
        "lessons": [
            ("Identify and scope a problem", [
                ("Industry problem selection", "Choose a specific, observable user problem rather than starting from a technology or a large feature list.", "Students miss event deadlines because dates and registration links are scattered."),
                ("Requirement analysis", "Requirements describe user goals and constraints; acceptance criteria make expected behavior testable.", "Given an open event, when a student registers, then they see a confirmation."),
                ("Scope definition", "An MVP is the smallest useful slice that can be completed and evaluated; defer lower-priority ideas explicitly.", "Build browse/detail/register first; defer analytics and notifications."),
                ("Design thinking and user journeys", "Understand users, map their steps and friction, prototype options, then use feedback to refine the problem.", "Sketch how a student discovers an event, checks seats, and confirms attendance."),
            ]),
            ("Design the solution", [
                ("User stories", "A user story states a goal from a user perspective; acceptance criteria define observable completion.", "As a student, I can see the start time before registering so I can decide to attend."),
                ("Wireframing", "A wireframe lays out content and interaction before visual polish, exposing missing states early.", "Sketch list, no-results, event detail, form error, and confirmation states."),
                ("Architecture design", "Architecture assigns responsibilities and data flow to components/services while making risks and trade-offs explicit.", "React owns presentation; Express enforces capacity; MongoDB persists registrations."),
                ("Database design", "Database design chooses fields, relationships, validation, and indexes based on domain rules and query patterns.", "A registration references an event and student; define uniqueness to prevent duplicate signup."),
            ]),
            ("Build, test and release", [
                ("Frontend development", "The frontend presents workflows, validates for usability, and reflects server outcomes.", "Disable duplicate submit while the registration request is pending."),
                ("Backend and database integration", "Backend routes enforce business rules and persist data through a defined data model.", "Check remaining capacity and write the registration as one server-side operation."),
                ("Authentication", "Authentication establishes user identity; include it only when the selected user workflow requires it.", "Require login to view personal registrations, not just because the stack includes JWT."),
                ("Functional testing and bug fixing", "Test acceptance criteria and edge cases; fix reproducible defects and rerun the relevant check.", "Verify full event, duplicate signup, unknown event, and normal registration."),
                ("Performance improvements", "Measure a real bottleneck before optimizing; preserve correctness while improving the measured path.", "Check API response time and payload size before adding a caching layer."),
                ("Production deployment", "A production deployment configures services, secrets, and network access, then verifies the live user path.", "Open the deployed site and complete a test registration using fictional data."),
            ]),
            ("Work as a team and present", [
                ("Sprint-based development", "A sprint is a short planning and feedback cycle with a goal, visible work, and review; it is not a promise of fixed velocity.", "Commit to a working registration slice by the next mentor checkpoint."),
                ("Peer review", "Code review checks correctness, clarity, risk, and tests while sharing context across the team.", "Reviewer asks how the capacity check prevents concurrent overbooking."),
                ("Technical demonstration", "A demo shows a real user workflow and its observable result, including meaningful failure behavior.", "Show an attendee registering, then show the full-event response."),
                ("Project documentation", "Documentation helps someone understand purpose, setup, architecture, verification, and known limits.", "A README includes prerequisites, environment variable names, run steps, and test command."),
                ("Presentation and walkthrough", "A concise presentation connects user need to design, implementation, evidence, trade-offs, and next steps.", "Each member explains their contribution and one decision they would revisit."),
            ]),
        ],
        "practice": "Use the paired capstone studio checklist across all 24 hours. Hold mentor demos at problem, design, working-slice, verification, and release checkpoints.",
    },
    {
        "id": "10", "title": "Interview Preparation & Placement Readiness",
        "duration": "6 hours", "outcome": "Solve a basic coding prompt aloud and present verifiable project and portfolio evidence.",
        "thread": "Practice explaining a capstone contribution honestly and clearly to a technical interviewer.",
        "lessons": [
            ("Coding assessment foundations", [
                ("Arrays and strings", "Arrays keep ordered items; strings represent text. Clarify indexing and empty-input behavior before solving.", "Scan registration IDs in an array to find a repeated value."),
                ("Hashing basics", "A hash map/set provides average constant-time lookup by key, using extra memory to trade for speed.", "Use a Set to remember IDs already encountered."),
                ("Sorting and searching", "Sorting orders values; searching locates a target. Binary search requires sorted input and halves the search range.", "Sort events by date for display; binary-search IDs only if the list is sorted."),
                ("Problem-solving strategy", "Restate the prompt, ask constraints, work examples, choose a simple approach, test edge cases, then discuss complexity.", "Trace `[4, 8, 4]` and `[]` before coding the duplicate-ID solution."),
            ]),
            ("Professional profile and communication", [
                ("Resume building", "A resume summarizes relevant evidence in concise, truthful, role-aligned language.", "Describe your API contribution with its tested behavior rather than claiming the whole team project."),
                ("GitHub portfolio", "A portfolio makes code and project outcomes inspectable through focused repositories, setup notes, and evidence.", "Add a README with a screenshot, run steps, tests, and your contribution."),
                ("LinkedIn profile", "A professional profile states skills and interests consistently with verifiable projects and experience.", "Link the deployed capstone and describe your role accurately."),
                ("Self introduction", "A concise introduction connects current study, relevant skill, evidence, and the role of interest.", "“I study computing, built a MERN event flow, and focused on server-side validation.”"),
                ("Group discussions", "Effective discussion combines clear points, active listening, relevant evidence, and space for others.", "Build on a classmate's point with a concrete example, then invite another view."),
                ("Presentation skills", "A technical presentation gives the audience a clear path through problem, solution, evidence, and limitations.", "Demo one working registration and explain one known limitation."),
            ]),
            ("Interview practice", [
                ("Technical interviews", "Technical interviews assess reasoning and communication as well as code; clarification and testing are part of the work.", "Ask whether duplicate IDs can appear before selecting the solution structure."),
                ("HR and behavioral interviews", "Behavioral responses use a real situation, your task, your actions, and a concrete result/lesson.", "Explain how you unblocked a team integration issue and what you learned."),
                ("Mock interviews", "Mock interviews build familiarity through realistic practice and specific, actionable feedback.", "Give one strength and one next step after a ten-minute peer interview."),
                ("Placement readiness portfolio", "A placement portfolio brings resume, profile, projects, documentation, and a concise pitch into a coherent evidence set.", "Check that the portfolio's live link works and each claim maps to project evidence."),
            ]),
        ],
        "practice": "Complete a timed but supportive coding interview, a 90-second project pitch, and one portfolio improvement. Keep feedback specific and actionable.",
    },
    {
        "id": "11", "title": "Advanced AI-Assisted Software Engineering",
        "duration": "10 hours", "outcome": "Design a bounded AI workflow and validate outputs for correctness, security, privacy, and value.",
        "thread": "Use an approved AI assistant on a small feature, then review and test its proposal before use.",
        "lessons": [
            ("AI-native engineering and collaboration", [
                ("Developer productivity evolution", "Tooling shifts effort over time; AI can accelerate some tasks but still creates review and integration work.", "Compare time to a tested event filter, not lines of code generated."),
                ("AI-native teams", "AI-native teams intentionally define where tools help, who reviews output, and how knowledge is retained.", "Agree that generated code needs a human owner and CI checks before merge."),
                ("Human-AI collaboration", "Humans set goals, constraints, and accountability; models generate suggestions that require verification.", "Ask AI for alternative event data models, then check each against required queries."),
                ("AI-augmented SDLC", "AI may assist requirements, design, coding, testing, or documentation while people remain responsible for decisions and releases.", "Generate test ideas, select relevant cases, implement and run them yourself."),
                ("Future trends", "AI capabilities and limitations change quickly; durable skills include domain understanding, evaluation, security, and communication.", "Reassess a tool against the course's privacy rules before adopting a new feature."),
            ]),
            ("Prompt and context engineering", [
                ("Structured prompt design", "A useful prompt states task, audience, verified context, constraints, output format, and how success will be checked.", "Ask for two API options, list assumptions, and include tests for full and open events."),
                ("Context window management", "A model can only use the context available within its limits; prioritize relevant facts and summarize with provenance.", "Provide the event schema and route contract instead of an unrelated full repository dump."),
                ("Knowledge grounding", "Grounding supplies trusted source material and asks the model to distinguish evidence from assumptions.", "Ask it to cite the supplied API contract for each generated field."),
                ("Retrieval-augmented workflows", "Retrieval finds relevant approved documents for a prompt; retrieved text still needs relevance and permission checks.", "Retrieve the current project README and route specification before proposing a change."),
                ("Reusable prompt frameworks", "A prompt framework standardizes inputs and review outputs while leaving room for task-specific context.", "Template: goal, files/context, constraints, output, assumptions, tests, risk review."),
            ]),
            ("Agents and bounded workflows", [
                ("AI agents", "An agent uses a model with tools and a loop to pursue a goal; tool access increases both capability and risk.", "A read-only agent can summarize API routes without permission to change files."),
                ("Task planning and delegation", "Break work into bounded tasks with explicit inputs, outputs, dependencies, owners, and stop conditions.", "One role drafts acceptance criteria; another reviews privacy; a person approves implementation."),
                ("Multi-agent collaboration", "Multiple agents can divide work but may duplicate assumptions or pass errors downstream; reconcile outputs centrally.", "Compare two architecture reviews against the same requirements before accepting either."),
                ("Human-in-the-loop", "A human approval gate is required before consequential actions such as changing protected data, credentials, or deployment.", "Require a person to inspect a patch and tests before merge or release."),
            ]),
            ("AI for architecture and system design", [
                ("Requirement analysis", "AI can draft questions and acceptance criteria from supplied needs; stakeholders must confirm meaning and priority.", "Ask what “event is available” means when capacity reaches zero."),
                ("User story generation", "AI can suggest story variants; teams verify that stories represent actual users and are testable.", "Review whether “organizer can edit event” names the permission and success condition."),
                ("Architecture review", "AI can surface possible risks and trade-offs but cannot know unstated constraints or guarantee correctness.", "Compare an embedded registration list with a separate collection against scale and query needs."),
                ("Pattern recommendations", "Patterns are context-dependent options, not prescriptions; assess operational and maintenance costs.", "Do not add a message queue for a small synchronous signup without a demonstrated need."),
                ("Technical decision support", "Use AI to enumerate options and questions; record the human decision, evidence, and unresolved uncertainty.", "Choose token expiry based on threat model and product needs, not a model's unsupported default."),
            ]),
            ("Governance, quality and adoption", [
                ("Risk management and hallucinations", "AI risk review considers impact, likelihood, uncertainty, data exposure, and ways to detect unsupported claims.", "Verify a claimed library method against official docs before shipping generated code."),
                ("Secure and responsible AI use", "Do not share secrets or private data with unapproved tools; check bias, privacy, licensing, and intended use.", "Use fictional event records in a prompt, never real attendee email addresses."),
                ("Output validation", "Validate generated artifacts with tests, source checks, threat review, and human domain judgment.", "Run a test proving an attendee cannot edit another person's registration."),
                ("Enterprise adoption", "Adoption needs approved tools, training, policy, support, and feedback loops aligned with organizational risk.", "Pilot an assistant on documentation before enabling repository write access."),
                ("Workflow standardization and knowledge repositories", "Shared playbooks and prompt libraries preserve approved practices and should be versioned and reviewed.", "Store a reviewed API-test prompt with its expected output and limitations."),
                ("Measure AI return on investment", "Measure time to verified outcome alongside quality, rework, defects, risk, and learning; compare equivalent tasks.", "Track review time and escaped defects before and after assistance, not token count."),
            ]),
        ],
        "practice": "Create the AI Engineering Playbook: reusable prompt, bounded workflow, governance checklist, validation record, and a quality-aware productivity metric.",
    },
]


def set_background(slide, color=PAPER):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, text, x, y, w, h, size, color=INK, bold=False,
             font="Aptos", align=PP_ALIGN.LEFT, margin=0):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_footer(slide, module_id, page, total):
    add_text(slide, f"CAMPUS EVENT DESK  /  MODULE {module_id}", 0.55, 7.14,
             7.0, 0.18, 8, MUTED, bold=True)
    add_text(slide, f"{page:02d}  /  {total:02d}", 11.45, 7.14,
             1.3, 0.18, 8, MUTED, align=PP_ALIGN.RIGHT)


def add_base_slide(presentation, title, eyebrow, module_id):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    set_background(slide)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.16), Inches(HEIGHT))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    add_text(slide, eyebrow.upper(), 0.62, 0.34, 11.9, 0.28, 10, TEAL, bold=True)
    add_text(slide, title, 0.62, 0.66, 12.0, 0.58, 25, INK, bold=True,
             font="Aptos Display")
    return slide


def add_cover(presentation, module):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    set_background(slide, INK)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(0.78),
                                    Inches(0.12), Inches(5.75))
    accent.fill.solid()
    accent.fill.fore_color.rgb = CORAL
    accent.line.fill.background()
    add_text(slide, f"MODULE {module['id']}  /  {module['duration'].upper()}",
             1.15, 1.02, 10.9, 0.35, 13, RGBColor(128, 214, 201), bold=True)
    add_text(slide, module["title"], 1.15, 1.66, 10.9, 1.65, 34, WHITE,
             bold=True, font="Aptos Display")
    add_text(slide, "A practical MERN course for engineering students",
             1.18, 3.52, 10.8, 0.42, 17, RGBColor(218, 231, 229))
    add_text(slide, f"BUILD THREAD   {module['thread']}", 1.18, 4.48, 10.8,
             0.88, 15, WHITE, bold=True)
    add_text(slide, "Explain it  /  demonstrate it  /  build it  /  verify it",
             1.18, 6.35, 10.8, 0.3, 11, RGBColor(175, 198, 197))
    return slide


def add_outcomes(presentation, module):
    slide = add_base_slide(presentation, "By the end of this module",
                           "Learning targets", module["id"])
    topics = [
        ("UNDERSTAND", module["outcome"],
         "Students should explain the concept in their own words, not just repeat a definition."),
        ("SEE IT", "Follow the concept through the Campus Event Desk example.",
         "The running scenario keeps technical ideas connected to a user and a real workflow."),
        ("BUILD IT", module["practice"],
         "The paired lab is the evidence of learning; the deck supports it rather than replacing it."),
    ]
    add_topic_cards(slide, topics)
    return slide


def add_topic_cards(slide, topics):
    start_y = 1.52
    available = 5.35
    gap = 0.14
    card_h = (available - gap * (len(topics) - 1)) / len(topics)
    for index, (topic, explanation, example) in enumerate(topics):
        y = start_y + index * (card_h + gap)
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(y),
                                       Inches(12.0), Inches(card_h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = RGBColor(221, 231, 227)
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(y),
                                        Inches(0.1), Inches(card_h))
        accent.fill.solid()
        accent.fill.fore_color.rgb = TEAL if index % 2 == 0 else CORAL
        accent.line.fill.background()
        top = y + 0.12
        add_text(slide, topic, 0.9, top, 11.3, 0.25, 14, TEAL, bold=True)
        add_text(slide, f"EXPLAIN  {explanation}", 0.9, top + 0.31, 11.25,
                 max(0.27, card_h * 0.36), 12, INK)
        add_text(slide, f"EXAMPLE  {example}", 0.9, top + card_h * 0.58,
                 11.25, max(0.27, card_h * 0.3), 11, MUTED)


def add_practice(presentation, module):
    slide = add_base_slide(presentation, "Practice, then check your understanding",
                           "Paired lab", module["id"])
    prompts = [
        ("1 / BUILD", module["practice"],
         "Use the matching hands-on instructions in training/hands-on."),
        ("2 / VERIFY", "Test a normal case, a boundary case, and a failure case.",
         "Write down what you observed; do not mark a feature complete from appearance alone."),
        ("3 / EXPLAIN", "What decision did you make, and what evidence supports it?",
         "A short partner walkthrough reveals whether the idea is understood and working."),
    ]
    add_topic_cards(slide, prompts)
    return slide


def add_exit(presentation, module):
    slide = add_base_slide(presentation, "Exit ticket", "Recall and transfer", module["id"])
    prompt = "In one minute: explain one concept, give a new example, and name one question you still have."
    add_text(slide, prompt, 0.86, 1.72, 11.2, 1.05, 25, INK, bold=True,
             font="Aptos Display")
    add_text(slide, "SAVE AS EVIDENCE", 0.88, 3.42, 3.5, 0.28, 11, CORAL, bold=True)
    add_text(slide, "One working artifact  +  one test/check  +  one sentence describing your decision",
             0.88, 3.84, 11.15, 0.72, 20, TEAL, bold=True)
    add_text(slide, "Next step: revisit the paired demo or lab before moving to the next module.",
             0.88, 5.45, 11.2, 0.5, 15, MUTED)
    return slide


def make_deck(module):
    presentation = Presentation()
    presentation.slide_width = Inches(WIDTH)
    presentation.slide_height = Inches(HEIGHT)
    add_cover(presentation, module)
    add_outcomes(presentation, module)

    for lesson_title, topics in module["lessons"]:
        chunk_count = (len(topics) + 2) // 3
        base_size, extra = divmod(len(topics), chunk_count)
        chunks = []
        start = 0
        for chunk_index in range(chunk_count):
            chunk_size = base_size + (1 if chunk_index < extra else 0)
            chunks.append(topics[start:start + chunk_size])
            start += chunk_size
        for chunk_index, chunk in enumerate(chunks, start=1):
            title = lesson_title if len(chunks) == 1 else f"{lesson_title}  /  Part {chunk_index}"
            slide = add_base_slide(presentation, title, "Concepts with examples", module["id"])
            add_topic_cards(slide, chunk)

    add_practice(presentation, module)
    add_exit(presentation, module)
    total = len(presentation.slides)
    for page, slide in enumerate(presentation.slides, start=1):
        if page != 1:
            add_footer(slide, module["id"], page, total)
    return presentation


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for module in MODULES:
        output_file = OUTPUT / f"module-{module['id']}.pptx"
        make_deck(module).save(output_file)
        print(f"Created {output_file.relative_to(ROOT)}")


if __name__ == "__main__":
    main()